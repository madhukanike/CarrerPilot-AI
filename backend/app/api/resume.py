from fastapi import APIRouter, UploadFile, File, HTTPException
from pathlib import Path
import io

from app.agents.career_agent import CareerAgent


router = APIRouter(prefix="/resume", tags=["Resume"])

agent = CareerAgent()


def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF using PyMuPDF.
    """
    try:
        import fitz

        document = fitz.open(
            stream=file_bytes,
            filetype="pdf"
        )

        pages = []

        for page in document:
            text = page.get_text("text")

            if text:
                pages.append(text)

        document.close()

        return "\n".join(pages).strip()

    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="PyMuPDF is not installed. Install it with: pip install pymupdf"
        )

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not extract text from PDF: {str(exc)}"
        )


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX file.
    """
    try:
        from docx import Document

        document = Document(
            io.BytesIO(file_bytes)
        )

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract text from tables.
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    paragraphs.append(" | ".join(row_text))

        return "\n".join(paragraphs).strip()

    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="python-docx is not installed. Install it with: pip install python-docx"
        )

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not extract text from DOCX: {str(exc)}"
        )


def extract_txt_text(file_bytes: bytes) -> str:
    """
    Extract text from a TXT file.
    """
    try:
        # UTF-8 first.
        return file_bytes.decode("utf-8").strip()

    except UnicodeDecodeError:
        try:
            # Fallback for files created with Windows encoding.
            return file_bytes.decode("latin-1").strip()

        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read TXT file: {str(exc)}"
            )


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """
    Detect the file type and extract the original resume text.
    """
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    if extension == ".txt":
        return extract_txt_text(file_bytes)

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Please upload PDF, DOCX, or TXT."
    )


@router.post("/upload")
async def upload_resume(file: UploadFile = File(...)):
    """
    Upload a resume, extract its original text,
    analyze it using AI, and return both the
    extracted text and AI analysis.
    """

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No filename was provided."
        )

    filename = file.filename
    extension = Path(filename).suffix.lower()

    allowed_extensions = {
        ".pdf",
        ".docx",
        ".txt",
    }

    if extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Please upload PDF, DOCX, or TXT."
        )

    try:
        file_bytes = await file.read()

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read uploaded file: {str(exc)}"
        )

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    # Extract the original resume text.
    resume_text = extract_resume_text(
        file_bytes,
        filename
    )

    if not resume_text:
        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text was found in the resume. "
                "If this is a scanned PDF, OCR support will be added later."
            )
        )

    # Protect the AI request from unnecessarily large files.
    # This does NOT change the text returned to the frontend.
    analysis_text = resume_text[:30000]

    try:
        analysis = agent.analyze_resume(
            analysis_text
        )

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Resume AI analysis failed: {str(exc)}"
        )

    return {
        "filename": filename,
        "resume_text": resume_text,
        "text": resume_text,
        "extracted_text": resume_text,
        "analysis": analysis,
    }